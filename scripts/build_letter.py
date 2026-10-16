"""Render a one-page cover letter (.docx, optional .pdf) from JSON.

Usage: python build_letter.py letter.json "Cover Letter - Company - Role.docx" [--pdf]
JSON: {"name", "contact": [...], "date", "recipient": [...], "salutation", "paragraphs": [...], "closing"}
Requires: pip install python-docx
"""
import argparse, json, os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor

sys.path.insert(0, os.path.dirname(__file__))
from docx_common import to_pdf


def build(L, out):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.4)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)

    def p(text, bold=False, size=11, after=6, color=None):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(after)
        r = para.add_run(text)
        r.bold, r.font.size = bold, Pt(size)
        if color:
            r.font.color.rgb = color
        return para

    p(L["name"], True, 16, 0, RGBColor(0x1F, 0x3A, 0x5F))
    if L.get("contact"):
        p("  |  ".join(L["contact"]), size=9.5, after=14)
    if L.get("date"):
        p(L["date"], after=10)
    for line in L.get("recipient", []):
        p(line, after=0)
    p("", after=6)
    p(L.get("salutation", "Dear Hiring Team,"), after=8)
    for para_text in L["paragraphs"]:
        p(para_text, after=8)
    p(L.get("closing", "Kind regards,"), after=2)
    p(L["name"], True)
    doc.save(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("json")
    ap.add_argument("out")
    ap.add_argument("--pdf", action="store_true")
    a = ap.parse_args()
    build(json.load(open(a.json, encoding="utf-8")), a.out)
    print("Wrote", a.out)
    if a.pdf:
        print(to_pdf(a.out))
