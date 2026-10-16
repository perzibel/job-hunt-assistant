"""Render a clean, ATS-friendly CV (.docx, optional .pdf) from JSON.

Usage: python build_cv.py cv.json "Out Name CV.docx" [--pdf]
Schema: see references/cv-tailoring.md. Requires: pip install python-docx
"""
import argparse, json, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

sys.path.insert(0, __import__("os").path.dirname(__file__))
from docx_common import to_pdf

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def para(doc, text="", size=10, bold=False, color=None, align=None, after=2, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(after), Pt(before)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size, r.bold = Pt(size), bold
        if color:
            r.font.color.rgb = color
    return p


def rule(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A5F")):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def build(cv, out):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)

    para(doc, cv["name"], 20, True, ACCENT, WD_ALIGN_PARAGRAPH.CENTER, 2)
    if cv.get("contact"):
        para(doc, "  |  ".join(cv["contact"]), 9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    if cv.get("summary"):
        para(doc, cv["summary"], 10, after=6)

    for sec in cv.get("sections", []):
        h = para(doc, sec["title"].upper(), 11, True, ACCENT, before=6, after=4)
        rule(h)
        kind = sec.get("type", "text")
        if kind == "text":
            para(doc, sec.get("text", ""), after=4)
        elif kind == "skills":
            for it in sec["items"]:
                p = para(doc, after=1)
                p.add_run(it["label"] + ": ").bold = True
                p.add_run(it["text"])
        elif kind == "jobs":
            for job in sec["items"]:
                p = para(doc, after=0, before=4)
                tab = p.paragraph_format.tab_stops
                tab.add_tab_stop(Cm(17.4), alignment=2)  # right-aligned dates
                p.add_run(job["title"]).bold = True
                org = ", ".join(x for x in (job.get("org"), job.get("location")) if x)
                if org:
                    p.add_run("  ·  " + org)
                if job.get("dates"):
                    r = p.add_run("\t" + job["dates"])
                    r.italic = True
                if job.get("intro"):
                    para(doc, job["intro"], 9.5, after=1).runs[0].italic = True
                for b in job.get("bullets", []):
                    bp = doc.add_paragraph(b, style="List Bullet")
                    bp.paragraph_format.space_after = Pt(1)
    doc.save(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("json")
    ap.add_argument("out")
    ap.add_argument("--pdf", action="store_true")
    a = ap.parse_args()
    build(json.load(open(a.json, encoding="utf-8")), a.out)
    print("Wrote", a.out)
    if a.pdf:
        print(to_pdf(a.out))
